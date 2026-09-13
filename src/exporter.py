"""File output for clean reviews and their optional analysis results.

The exporter only writes the supplied records; filtering belongs to the service.
Files are published after writing succeeds so a failed export cannot replace a
previous result. CSV escapes formula-like text with an apostrophe for spreadsheet
safety. JSONL and Excel retain the original strings (Excel cells are plain text).
"""

from __future__ import annotations

import csv
import json
import os
import tempfile
from contextlib import suppress
from pathlib import Path
from typing import Any, Dict, Sequence

from src.config import resolve_project_path
from src.errors import OutputError, ValidationError
from src.models import (
    ExportFormat,
    ExportResult,
    OutputArtifact,
    OutputKind,
    ReviewDetail,
)


EXPORT_FIELDS = (
    "id", "source_review_id", "product_name", "review_date", "rating",
    "review_text", "cleaned_at", "sentiment", "confidence", "summary",
    "keywords", "analyzed_at", "provider", "model", "prompt_version",
)
_SUFFIXES = {
    ExportFormat.CSV: ".csv",
    ExportFormat.JSONL: ".jsonl",
    ExportFormat.EXCEL: ".xlsx",
}
_EXCEL_MAX_ROWS = 1_048_576
_EXCEL_MAX_TEXT_LENGTH = 32_767


def _flatten(detail: ReviewDetail) -> Dict[str, Any]:
    review, analysis = detail.review, detail.analysis
    return {
        "id": review.id,
        "source_review_id": review.source_review_id,
        "product_name": review.product_name,
        "review_date": review.review_date.isoformat(),
        "rating": review.rating,
        "review_text": review.review_text,
        "cleaned_at": review.cleaned_at.isoformat().replace("+00:00", "Z"),
        "sentiment": analysis.sentiment.value if analysis else None,
        "confidence": analysis.confidence if analysis else None,
        "summary": analysis.summary if analysis else None,
        "keywords": list(analysis.keywords) if analysis else None,
        "analyzed_at": (
            analysis.analyzed_at.isoformat().replace("+00:00", "Z")
            if analysis else None
        ),
        "provider": analysis.provider if analysis else None,
        "model": analysis.model if analysis else None,
        "prompt_version": analysis.prompt_version if analysis else None,
    }


def _table_value(value: Any) -> Any:
    if isinstance(value, list):
        return json.dumps(value, ensure_ascii=False)
    return value


def _csv_value(value: Any) -> Any:
    value = _table_value(value)
    # CSV cannot mark a cell as text. Prefix dangerous strings, including ones
    # hidden behind whitespace; numeric rating/confidence values remain numeric.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


class FileReviewExporter:
    """Implement the ReviewExporter boundary using CSV, JSONL, or Excel files."""

    def export_reviews(
        self,
        reviews: Sequence[ReviewDetail],
        output: Path,
        *,
        export_format: ExportFormat,
        force: bool = False,
    ) -> ExportResult:
        if not isinstance(export_format, ExportFormat):
            raise ValidationError("export_format must be an ExportFormat value")
        if not isinstance(force, bool):
            raise ValidationError("force must be a boolean")
        if not isinstance(output, Path):
            raise ValidationError("output must be a Path")
        if output.suffix and output.suffix.lower() != _SUFFIXES[export_format]:
            raise ValidationError("출력 파일 확장자가 내보내기 형식과 일치하지 않습니다.")
        if any(not isinstance(review, ReviewDetail) for review in reviews):
            raise ValidationError("reviews must contain ReviewDetail values")

        temporary: Path | None = None
        try:
            # Resolve parent links without following a newly created final link.
            target = resolve_project_path(output.parent) / output.name
            if target.is_symlink():
                raise OutputError("심볼릭 링크에는 내보낼 수 없습니다.")
            if target.exists() and not force:
                raise OutputError("출력 파일이 이미 있습니다. 덮어쓰려면 --force를 지정하세요.")
            target.parent.mkdir(parents=True, exist_ok=True)
            descriptor, name = tempfile.mkstemp(
                prefix=".review-export-", suffix=_SUFFIXES[export_format],
                dir=target.parent,
            )
            temporary = Path(name)
            os.close(descriptor)

            if export_format is ExportFormat.CSV:
                self._write_csv(reviews, temporary)
            elif export_format is ExportFormat.JSONL:
                self._write_jsonl(reviews, temporary)
            else:
                self._write_excel(reviews, temporary)

            if target.is_symlink():
                raise OutputError("심볼릭 링크에는 내보낼 수 없습니다.")
            if force:
                os.replace(temporary, target)
            else:
                # Unlike exists() + replace(), link atomically refuses an existing
                # destination even when another writer creates it during export.
                try:
                    os.link(temporary, target)
                except FileExistsError:
                    raise OutputError(
                        "출력 파일이 이미 있습니다. 덮어쓰려면 --force를 지정하세요."
                    ) from None
            return ExportResult(
                artifact=OutputArtifact(OutputKind.EXPORT, target, export_format.value),
                row_count=len(reviews),
            )
        except OutputError:
            raise
        except Exception:
            # Third-party writers may put the full cell content in an exception.
            # Surface only a safe application error to the CLI or caller.
            raise OutputError("내보내기 파일을 생성할 수 없습니다.") from None
        finally:
            if temporary is not None:
                with suppress(OSError):
                    temporary.unlink(missing_ok=True)

    @staticmethod
    def _write_csv(reviews: Sequence[ReviewDetail], output: Path) -> None:
        with output.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=EXPORT_FIELDS)
            writer.writeheader()
            for review in reviews:
                writer.writerow({key: _csv_value(value) for key, value in _flatten(review).items()})

    @staticmethod
    def _write_jsonl(reviews: Sequence[ReviewDetail], output: Path) -> None:
        with output.open("w", encoding="utf-8", newline="\n") as stream:
            for review in reviews:
                stream.write(json.dumps(_flatten(review), ensure_ascii=False, allow_nan=False))
                stream.write("\n")

    @staticmethod
    def _write_excel(reviews: Sequence[ReviewDetail], output: Path) -> None:
        try:
            from openpyxl import Workbook
        except ImportError:
            raise OutputError("Excel 내보내기에는 openpyxl 패키지가 필요합니다.") from None

        if len(reviews) + 1 > _EXCEL_MAX_ROWS:
            raise OutputError("Excel 행 수 제한을 초과했습니다. CSV 또는 JSONL을 사용하세요.")
        workbook = Workbook()
        try:
            sheet = workbook.active
            sheet.title = "Reviews"
            sheet.append(EXPORT_FIELDS)
            sheet.freeze_panes = "A2"
            for row_index, detail in enumerate(reviews, start=2):
                for column, value in enumerate(_flatten(detail).values(), start=1):
                    value = _table_value(value)
                    if isinstance(value, str) and len(value) > _EXCEL_MAX_TEXT_LENGTH:
                        # openpyxl silently slices long strings; reject first so
                        # that exported review content is never silently lost.
                        raise OutputError(
                            "Excel 셀의 문자열 길이 제한을 초과했습니다. "
                            "CSV 또는 JSONL을 사용하세요."
                        )
                    cell = sheet.cell(row=row_index, column=column, value=value)
                    if isinstance(value, str):
                        cell.data_type = "s"
            sheet.auto_filter.ref = f"A1:O{len(reviews) + 1}"
            for column, field_name in enumerate(EXPORT_FIELDS, start=1):
                width = 60 if field_name in ("review_text", "summary") else 24
                if field_name in ("id", "rating", "confidence"):
                    width = 12
                sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
            workbook.save(output)
        finally:
            workbook.close()


__all__ = ["EXPORT_FIELDS", "FileReviewExporter"]
