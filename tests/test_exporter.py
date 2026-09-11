"""Reread exported files and verify safe publication at the output boundary."""

import builtins
import contextlib
import csv
import io
import json
import tempfile
import unittest
from datetime import date, datetime, timezone
from pathlib import Path
from unittest.mock import patch

from src.errors import OutputError, ValidationError
from src.exporter import EXPORT_FIELDS, FileReviewExporter
from src.models import (
    AnalysisResult, CleanReview, ExportFormat, OutputKind, ReviewDetail, Sentiment,
)
from src.services import ReviewExporter


class FileReviewExporterTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name).resolve()
        self.exporter = FileReviewExporter()
        self.reviews = [self._detail(1, analyzed=True), self._detail(2)]

    @staticmethod
    def _detail(review_id, *, analyzed=False, text='좋아요, "편해요"\n다음 줄 😀'):
        review = CleanReview(
            id=review_id, source_review_id="0007" if analyzed else None,
            product_name="무선 이어폰 🎧", review_date=date(2026, 9, 11),
            rating=5, review_text=text,
            cleaned_at=datetime(2026, 9, 11, 1, 2, 3, 123456, tzinfo=timezone.utc),
        )
        analysis = AnalysisResult(
            review_id=review_id, sentiment=Sentiment.POSITIVE, confidence=0.95,
            analyzed_at=datetime(2026, 9, 11, 2, 3, 4, tzinfo=timezone.utc),
            provider="example", model="test-model", summary="착용감이 좋아요.\n만족 😀",
            keywords=["착용감", '한글, "키워드"', "😀"], prompt_version="v1",
        ) if analyzed else None
        return ReviewDetail(review, analysis)

    def _export(self, format, *, name=None, reviews=None, force=False):
        suffix = {ExportFormat.CSV: "csv", ExportFormat.JSONL: "jsonl", ExportFormat.EXCEL: "xlsx"}[format]
        target = self.root / (name or f"nested/reviews.{suffix}")
        result = self.exporter.export_reviews(
            self.reviews if reviews is None else reviews,
            target, export_format=format, force=force,
        )
        self.assertEqual(result.artifact.path, target.resolve())
        self.assertEqual(result.artifact.kind, OutputKind.EXPORT)
        self.assertEqual(result.artifact.format, format.value)
        self.assertEqual(result.row_count, len(self.reviews if reviews is None else reviews))
        return target

    def _assert_no_temporary_files(self):
        self.assertEqual(list(self.root.rglob(".review-export-*")), [])

    def test_implements_protocol(self):
        self.assertIsInstance(self.exporter, ReviewExporter)

    def test_csv_roundtrip_bom_flat_fields_and_multiline_text(self):
        target = self._export(ExportFormat.CSV)
        self.assertTrue(target.read_bytes().startswith(b"\xef\xbb\xbf"))
        with target.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            self.assertEqual(reader.fieldnames, list(EXPORT_FIELDS))
            rows = list(reader)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[0]["source_review_id"], "0007")
        self.assertEqual(rows[0]["product_name"], self.reviews[0].review.product_name)
        self.assertEqual(rows[0]["review_text"], self.reviews[0].review.review_text)
        self.assertEqual(rows[0]["review_date"], "2026-09-11")
        self.assertEqual(rows[0]["cleaned_at"], "2026-09-11T01:02:03.123456Z")
        self.assertEqual(rows[0]["analyzed_at"], "2026-09-11T02:03:04Z")
        self.assertEqual(rows[0]["summary"], self.reviews[0].analysis.summary)
        self.assertEqual(json.loads(rows[0]["keywords"]), self.reviews[0].analysis.keywords)
        self.assertEqual(rows[0]["confidence"], "0.95")
        for key in ("source_review_id", "sentiment", "confidence", "summary", "keywords", "analyzed_at", "provider", "model", "prompt_version"):
            self.assertEqual(rows[1][key], "")
        self._assert_no_temporary_files()

    def test_jsonl_roundtrip_utf8_dates_nulls_arrays_and_native_numbers(self):
        target = self._export(ExportFormat.JSONL)
        self.assertFalse(target.read_bytes().startswith(b"\xef\xbb\xbf"))
        text = target.read_text(encoding="utf-8")
        self.assertIn("무선 이어폰 🎧", text)
        rows = [json.loads(line) for line in text.splitlines()]
        self.assertEqual(len(rows), 2)
        self.assertEqual(list(rows[0]), list(EXPORT_FIELDS))
        self.assertEqual(rows[0]["id"], 1)
        self.assertEqual(rows[0]["rating"], 5)
        self.assertEqual(rows[0]["confidence"], 0.95)
        self.assertEqual(rows[0]["review_text"], self.reviews[0].review.review_text)
        self.assertEqual(rows[0]["keywords"], self.reviews[0].analysis.keywords)
        self.assertEqual(rows[0]["review_date"], "2026-09-11")
        self.assertEqual(rows[0]["cleaned_at"], "2026-09-11T01:02:03.123456Z")
        self.assertEqual(rows[0]["analyzed_at"], "2026-09-11T02:03:04Z")
        for key in ("source_review_id", "sentiment", "confidence", "summary", "keywords", "analyzed_at", "provider", "model", "prompt_version"):
            self.assertIsNone(rows[1][key])
        self._assert_no_temporary_files()

    def test_excel_roundtrip_values_and_navigation(self):
        from openpyxl import load_workbook

        target = self._export(ExportFormat.EXCEL)
        workbook = load_workbook(target)
        try:
            sheet = workbook.active
            rows = list(sheet.values)
            self.assertEqual(rows[0], EXPORT_FIELDS)
            self.assertEqual(len(rows), 3)
            first = dict(zip(EXPORT_FIELDS, rows[1]))
            second = dict(zip(EXPORT_FIELDS, rows[2]))
            self.assertEqual(first["source_review_id"], "0007")
            self.assertEqual(first["product_name"], self.reviews[0].review.product_name)
            self.assertEqual(first["review_text"], self.reviews[0].review.review_text)
            self.assertEqual(first["confidence"], 0.95)
            self.assertEqual(first["rating"], 5)
            self.assertEqual(first["review_date"], "2026-09-11")
            self.assertEqual(first["cleaned_at"], "2026-09-11T01:02:03.123456Z")
            self.assertEqual(first["analyzed_at"], "2026-09-11T02:03:04Z")
            self.assertEqual(json.loads(first["keywords"]), self.reviews[0].analysis.keywords)
            self.assertIsNone(second["sentiment"])
            self.assertIsNone(second["keywords"])
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:O3")
        finally:
            workbook.close()

    def test_empty_formats_have_only_headers_or_no_jsonl_records(self):
        from openpyxl import load_workbook

        csv_path = self._export(ExportFormat.CSV, reviews=[])
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            self.assertEqual(list(csv.reader(stream)), [list(EXPORT_FIELDS)])
        jsonl_path = self._export(ExportFormat.JSONL, reviews=[])
        self.assertEqual(jsonl_path.read_bytes(), b"")
        excel_path = self._export(ExportFormat.EXCEL, reviews=[])
        workbook = load_workbook(excel_path)
        try:
            self.assertEqual(list(workbook.active.values), [EXPORT_FIELDS])
            self.assertEqual(workbook.active.auto_filter.ref, "A1:O1")
        finally:
            workbook.close()

    def test_empty_keyword_list_is_distinct_from_missing_analysis(self):
        self.reviews[0].analysis.keywords = []
        target = self._export(ExportFormat.JSONL)
        rows = [json.loads(line) for line in target.read_text().splitlines()]
        self.assertEqual(rows[0]["keywords"], [])
        self.assertIsNone(rows[1]["keywords"])

    def test_all_formats_preserve_existing_files_unless_force(self):
        for format in ExportFormat:
            with self.subTest(format=format):
                target = self._export(format)
                target.write_bytes(b"original result")
                with self.assertRaises(OutputError):
                    self.exporter.export_reviews(self.reviews, target, export_format=format)
                self.assertEqual(target.read_bytes(), b"original result")
                self._export(format, force=True)
                self.assertNotEqual(target.read_bytes(), b"original result")
        self._assert_no_temporary_files()

    def test_writer_failure_preserves_previous_file_and_cleans_partial_output(self):
        for has_original in (False, True):
            with self.subTest(has_original=has_original):
                target = self.root / f"failure-{has_original}.csv"
                if has_original:
                    target.write_bytes(b"original result")

                def fail(reviews, path):
                    path.write_text("partial", encoding="utf-8")
                    raise RuntimeError("sensitive raw review text")

                captured = io.StringIO()
                with patch.object(self.exporter, "_write_csv", side_effect=fail), contextlib.redirect_stdout(captured), contextlib.redirect_stderr(captured):
                    with self.assertRaises(OutputError) as error:
                        self.exporter.export_reviews(self.reviews, target, export_format=ExportFormat.CSV, force=True)
                self.assertNotIn("sensitive", str(error.exception))
                self.assertEqual(captured.getvalue(), "")
                if has_original:
                    self.assertEqual(target.read_bytes(), b"original result")
                else:
                    self.assertFalse(target.exists())
                self._assert_no_temporary_files()

    def test_file_created_by_concurrent_writer_is_not_overwritten(self):
        target = self.root / "race.csv"
        original_write = self.exporter._write_csv

        def race(reviews, path):
            original_write(reviews, path)
            target.write_bytes(b"concurrent export")

        with patch.object(self.exporter, "_write_csv", side_effect=race):
            with self.assertRaises(OutputError):
                self.exporter.export_reviews(self.reviews, target, export_format=ExportFormat.CSV)
        self.assertEqual(target.read_bytes(), b"concurrent export")
        self._assert_no_temporary_files()

    def test_replace_failure_keeps_existing_file(self):
        target = self.root / "failure.csv"
        target.write_bytes(b"original result")
        with patch("src.exporter.os.replace", side_effect=PermissionError("denied")):
            with self.assertRaises(OutputError):
                self.exporter.export_reviews(self.reviews, target, export_format=ExportFormat.CSV, force=True)
        self.assertEqual(target.read_bytes(), b"original result")
        self._assert_no_temporary_files()

    def test_csv_formula_prefixes_are_escaped_but_jsonl_preserves_original(self):
        texts = ["=1+1", "+SUM(A1)", "-1+1", "@SUM(A1)", " \t=1+1", "\n+1+1", "정상 리뷰"]
        reviews = [self._detail(index, text=text) for index, text in enumerate(texts, start=1)]
        reviews[0].review.product_name = "=HYPERLINK(\"example\")"
        csv_path = self._export(ExportFormat.CSV, reviews=reviews)
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual([row["review_text"] for row in rows], ["'" + text for text in texts[:-1]] + [texts[-1]])
        self.assertEqual(rows[0]["product_name"], "'" + reviews[0].review.product_name)
        jsonl_path = self._export(ExportFormat.JSONL, reviews=reviews)
        rows = [json.loads(line) for line in jsonl_path.read_text().splitlines()]
        self.assertEqual([row["review_text"] for row in rows], texts)

    def test_excel_formula_like_and_error_like_text_remain_plain_text(self):
        from openpyxl import load_workbook

        texts = ["=SUM(A1:A2)", " +SUM(A1)", "#N/A", "@SUM(A1)"]
        reviews = [self._detail(index, text=text) for index, text in enumerate(texts, start=1)]
        target = self._export(ExportFormat.EXCEL, reviews=reviews)
        workbook = load_workbook(target, data_only=False)
        try:
            cells = [workbook.active.cell(row, 6) for row in range(2, len(texts) + 2)]
            self.assertEqual([cell.value for cell in cells], texts)
            self.assertEqual([cell.data_type for cell in cells], ["s"] * len(texts))
        finally:
            workbook.close()

    def test_excel_rejects_long_or_illegal_text_without_destroying_original(self):
        for text in ("가" * 32768, "민감한 내용\x00"):
            with self.subTest(length=len(text)):
                target = self.root / "long.xlsx"
                target.write_bytes(b"original workbook")
                with self.assertRaises(OutputError) as error:
                    self.exporter.export_reviews([self._detail(1, text=text)], target, export_format=ExportFormat.EXCEL, force=True)
                self.assertNotIn("민감한", str(error.exception))
                self.assertEqual(target.read_bytes(), b"original workbook")
                self._assert_no_temporary_files()

    def test_excel_row_limit_includes_header(self):
        with patch("src.exporter._EXCEL_MAX_ROWS", 2):
            with self.assertRaises(OutputError):
                self._export(ExportFormat.EXCEL)
            self._export(ExportFormat.EXCEL, reviews=self.reviews[:1])
        self._assert_no_temporary_files()

    def test_excel_dependency_is_optional_for_csv_and_jsonl(self):
        original_import = builtins.__import__

        def without_openpyxl(name, *args, **kwargs):
            if name == "openpyxl" or name.startswith("openpyxl."):
                raise ModuleNotFoundError("openpyxl is unavailable")
            return original_import(name, *args, **kwargs)

        with patch("builtins.__import__", side_effect=without_openpyxl):
            self._export(ExportFormat.CSV)
            self._export(ExportFormat.JSONL)
            with self.assertRaisesRegex(OutputError, "openpyxl"):
                self._export(ExportFormat.EXCEL)
        self._assert_no_temporary_files()

    def test_rejects_suffix_mismatch_and_invalid_options(self):
        cases = (
            (self.root / "reviews.xlsx", ExportFormat.CSV, False),
            (self.root / "reviews.csv", ExportFormat.EXCEL, False),
            (self.root / "reviews.csv", "csv", False),
            (self.root / "reviews.csv", ExportFormat.CSV, 1),
            ("reviews.csv", ExportFormat.CSV, False),
        )
        for target, format, force in cases:
            with self.subTest(target=target, format=format, force=force):
                with self.assertRaises(ValidationError):
                    self.exporter.export_reviews(self.reviews, target, export_format=format, force=force)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_suffixless_and_uppercase_names_are_supported(self):
        self._export(ExportFormat.CSV, name="no-extension")
        self._export(ExportFormat.EXCEL, name="REVIEWS.XLSX")

    def test_relative_paths_use_project_root(self):
        with patch("src.config.PROJECT_ROOT", self.root):
            result = self.exporter.export_reviews(self.reviews, Path("nested/relative.csv"), export_format=ExportFormat.CSV)
        self.assertEqual(result.artifact.path, self.root / "nested/relative.csv")
        self.assertTrue(result.artifact.path.exists())

    def test_symlink_destinations_are_rejected_even_with_force(self):
        destination = self.root / "original.csv"
        destination.write_bytes(b"original result")
        link = self.root / "link.csv"
        link.symlink_to(destination)
        for force in (False, True):
            with self.subTest(force=force):
                with self.assertRaises(OutputError):
                    self.exporter.export_reviews(self.reviews, link, export_format=ExportFormat.CSV, force=force)
                self.assertEqual(destination.read_bytes(), b"original result")
                self.assertTrue(link.is_symlink())
        broken_link = self.root / "broken.csv"
        broken_link.symlink_to(self.root / "absent.csv")
        with self.assertRaises(OutputError):
            self.exporter.export_reviews(self.reviews, broken_link, export_format=ExportFormat.CSV, force=True)
        self.assertFalse((self.root / "absent.csv").exists())

    def test_parent_file_failure_is_output_error(self):
        parent = self.root / "file"
        parent.write_text("file", encoding="utf-8")
        with self.assertRaises(OutputError):
            self.exporter.export_reviews(self.reviews, parent / "reviews.csv", export_format=ExportFormat.CSV)
        self.assertEqual(parent.read_text(), "file")
        self._assert_no_temporary_files()


if __name__ == "__main__":
    unittest.main()
