"""Absent optional metadata survives persistence, analysis and all consumers."""
import csv
import json
import sqlite3
import tempfile
import unittest
from dataclasses import replace
from datetime import date, datetime, timezone
from pathlib import Path
from unittest.mock import Mock

from openpyxl import Workbook, load_workbook
from src import cleaner, collector
from src.analyzer import SingleReviewAnalyzer
from src.ai_provider import ProviderResponse
from src.clean_service import CleanService
from src.comparison import ComparisonRequest, ComparisonService
from src.dashboard_service import DashboardService
from src.errors import ValidationError
from src.exporter import FileReviewExporter
from src.insight_batching import group_evidence
from src.insight_extractor import _matches
from src.models import (
    AnalysisOptions, AnalysisResult, CleaningOptions, CleanRequest, CleanReview,
    DashboardRequest, DuplicatePolicy, ExportFormat, InsightResult, RawReview,
    ReportFormat, ReviewDetail, ReviewFilter, ReviewQuery, Sentiment, SortField, SortOrder,
)
from src.query_output import format_review_list, format_review_detail, format_insight_result
from src.reporter import FileReportGenerator
from src.sqlite_repository import SQLiteReviewRepository
from src.visualizer import DashboardVisualizer
from src.web_dashboard import public_review


class OptionalReviewTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name).resolve()
        self.database = self.root / 'reviews.sqlite'
        self.repo = SQLiteReviewRepository(self.database)
        self.addCleanup(self.repo.close)
        self.now = datetime(2026, 9, 23, tzinfo=timezone.utc)
        self.clean = CleanService(self.repo, cleaner)

    def clean_rows(self, rows):
        self.repo.save_raw_reviews(rows, DuplicatePolicy.SKIP)
        return self.clean.clean_reviews(CleanRequest(CleaningOptions(DuplicatePolicy.UPSERT, 3)))

    def test_missing_and_blank_values_are_distinct_from_invalid_values(self):
        result = self.clean_rows([
            RawReview(source_review_id='missing', review_text='본문만 있는 유효한 리뷰'),
            RawReview(source_review_id='blank', product_name='  ', review_date='', rating=' ', review_text='빈 선택 항목 리뷰'),
            RawReview(source_review_id='bad-date', review_date='2026-02-30', review_text='잘못된 날짜 리뷰'),
            RawReview(source_review_id='bad-rating', rating=0, review_text='잘못된 별점 리뷰'),
            RawReview(source_review_id='empty', review_text=' '),
        ])
        self.assertEqual((result.succeeded, result.rejected), (2, 3))
        for review in result.reviews:
            self.assertEqual((review.product_name, review.review_date, review.rating), (None, None, None))
        self.assertEqual({error.code for error in result.errors},
                         {'INVALID_REVIEW_DATE', 'INVALID_RATING', 'MISSING_REVIEW_TEXT'})
        with sqlite3.connect(self.database) as connection:
            self.assertEqual(connection.execute('SELECT product_name,review_date,rating FROM clean_reviews').fetchall(),
                             [(None, None, None), (None, None, None)])

    def test_csv_and_excel_with_only_review_text_column(self):
        csv_path = self.root / 'input.csv'
        csv_path.write_text('review_text\n본문만 있는 리뷰입니다\n', encoding='utf-8-sig')
        workbook = Workbook()
        workbook.active.append(['review_text'])
        workbook.active.append(['Only the review body is provided.'])
        excel_path = self.root / 'input.xlsx'
        workbook.save(excel_path)
        workbook.close()
        for path in (csv_path, excel_path):
            with self.subTest(path=path):
                loaded = collector.load_reviews(path)
                self.repo.save_raw_reviews(loaded, DuplicatePolicy.SKIP)
        result = self.clean.clean_reviews(CleanRequest(CleaningOptions(DuplicatePolicy.SKIP, 3)))
        self.assertEqual(result.succeeded, 2)
        self.assertEqual(self.repo.get_statistics().total_reviews, 2)

    def seed_mixed(self):
        result = self.clean_rows([
            RawReview(source_review_id='1', review_text='배송이 늦었습니다'),
            RawReview(source_review_id='2', product_name='제품', rating=5,
                      review_date='2026-09-22', review_text='상품에 만족합니다'),
            RawReview(source_review_id='3', product_name='제품', rating=1, review_text='분석 실패 리뷰'),
            RawReview(source_review_id='4', product_name='제품', review_date='2026-09-22', review_text='아직 분석하지 않은 리뷰'),
        ])
        for review, sentiment in zip(result.reviews, (Sentiment.NEGATIVE, Sentiment.POSITIVE)):
            self.repo.save_analysis(AnalysisResult(review.id, sentiment, .9, self.now, 'fake', 'offline', keywords=['배송']))
        self.repo.mark_analysis_failed(3, 'test failure')

    def test_statistics_use_available_values_without_inventing_dates_or_ratings(self):
        self.seed_mixed()
        stats = self.repo.get_statistics()
        self.assertEqual((stats.total_reviews, stats.analyzed_reviews, stats.unanalyzed_reviews, stats.failed_reviews), (4, 2, 1, 1))
        self.assertEqual(stats.average_rating, 3)
        self.assertEqual(stats.sentiment_ratios[Sentiment.NEGATIVE], .5)
        self.assertEqual(sum(sum(row.values()) for row in stats.daily_sentiment_counts.values()), 1)
        self.assertEqual(sum(sum(row.values()) for row in stats.rating_sentiment_matrix.values()), 1)
        self.assertEqual(stats.top_negative_keywords[0].keyword, '배송')

    def test_filters_exclude_missing_values_and_sort_them_last_in_both_directions(self):
        self.seed_mixed()
        for filters, count in ((ReviewFilter(product_name='제품'), 3), (ReviewFilter(rating_min=1), 2),
                               (ReviewFilter(date_to=date(2026, 9, 22)), 2)):
            with self.subTest(filters=filters):
                self.assertEqual(self.repo.get_statistics(filters).total_reviews, count)
                self.assertFalse(_matches(self.repo.get_review(1), filters))
        self.assertTrue(_matches(self.repo.get_review(1), ReviewFilter(sentiment=Sentiment.NEGATIVE)))
        for sort, attribute in ((SortField.DATE, 'review_date'), (SortField.RATING, 'rating')):
            for order in SortOrder:
                page = self.repo.list_reviews(ReviewQuery(sort=sort, order=order))
                values = [getattr(d.review, attribute) for d in page.items]
                self.assertTrue(all(v is not None for v in values[:2]))
                self.assertEqual(values[2:], [None, None])

    def test_product_and_category_comparisons_keep_unnamed_reviews(self):
        self.seed_mixed()
        result = ComparisonService(self.repo).compare(ComparisonRequest())
        unknown = next(group for group in result.groups if group.name is None)
        self.assertEqual(unknown.label, '[제품명 없음]')
        self.assertEqual((unknown.statistics.total_reviews, unknown.product_count), (1, 0))
        category = ComparisonService(self.repo).compare(ComparisonRequest(group_by='category'))
        self.assertEqual(category.groups[0].statistics.total_reviews, 4)
        self.assertEqual(category.groups[0].product_count, 1)

    def test_query_web_and_export_outputs_preserve_null_or_show_missing(self):
        self.clean_rows([RawReview(review_text='본문만 있는 리뷰입니다')])
        detail = self.repo.get_review(1)
        text = format_review_detail(detail) + format_review_list(self.repo.list_reviews(ReviewQuery()))
        for expected in ('제품명 없음', '날짜 없음', 'N/A'):
            self.assertIn(expected, text)
        self.assertNotIn('None', text)
        self.assertIsNone(public_review(detail)['review_date'])
        exporter = FileReviewExporter()
        for fmt, suffix in ((ExportFormat.CSV, 'csv'), (ExportFormat.JSONL, 'jsonl'), (ExportFormat.EXCEL, 'xlsx')):
            path = self.root / f'reviews.{suffix}'
            exporter.export_reviews([detail], path, export_format=fmt)
        data = json.loads((self.root / 'reviews.jsonl').read_text())
        self.assertEqual([data[k] for k in ('product_name', 'review_date', 'rating')], [None] * 3)
        with (self.root / 'reviews.csv').open(encoding='utf-8-sig', newline='') as stream:
            row = next(csv.DictReader(stream))
        self.assertEqual([row[k] for k in ('product_name', 'review_date', 'rating')], [''] * 3)
        workbook = load_workbook(self.root / 'reviews.xlsx')
        self.addCleanup(workbook.close)
        self.assertEqual([workbook.active.cell(2, column).value for column in (3, 4, 5)], [None] * 3)

    def test_analysis_and_insight_evidence_support_unknown_product(self):
        review = self.clean_rows([RawReview(review_text='배송이 늦었습니다')]).reviews[0]
        provider = Mock()
        provider.complete.return_value = ProviderResponse(json.dumps({'sentiment': 'negative', 'confidence': .9,
            'summary': '배송이 늦었습니다.', 'keywords': ['배송']}), 'offline')
        analysis = SingleReviewAnalyzer(provider).analyze_review(review, AnalysisOptions('fake', 'offline', 30, 0))
        payload = json.loads(provider.complete.call_args.args[0][1]['content'])
        self.assertIsNone(payload['product_name'])
        self.assertIsNone(payload['rating'])
        detail = ReviewDetail(review, analysis)
        evidence = [{'complaints': [{'label': '배송 지연', 'quote': '배송이 늦었습니다'}], 'praises': []}]
        groups = group_evidence(evidence * 2, [detail, ReviewDetail(replace(review, id=2, product_name='제품'), replace(analysis, review_id=2))])
        insight = InsightResult(ReviewFilter(), 2, self.now, evidence_groups=groups)
        self.assertIn('제품명 없음', format_insight_result(insight))
        report = FileReportGenerator().generate_report(self.repo.get_statistics(), insight, self.root / 'insight.md', report_format=ReportFormat.MARKDOWN)
        self.assertIn('제품명 없음', report.path.read_text())

    def test_png_html_and_report_can_be_generated_with_only_body_and_analysis(self):
        review = self.clean_rows([RawReview(review_text='사용하기 편리합니다')]).reviews[0]
        self.repo.save_analysis(AnalysisResult(review.id, Sentiment.POSITIVE, .9, self.now, 'fake', 'offline'))
        result = DashboardService(self.repo, DashboardVisualizer(), FileReportGenerator(), dpi=50).create_dashboard(
            DashboardRequest(ReviewFilter(), self.root / 'dashboard', ReportFormat.MARKDOWN, generate_html=True))
        self.assertEqual(len(result.artifacts), 3)
        for artifact in result.artifacts:
            self.assertTrue(artifact.path.is_file())
        report = next(a.path for a in result.artifacts if a.format == 'md').read_text()
        self.assertIn('N/A', report)
        self.assertIn('작성일이 없는 리뷰', report)

    def test_optional_model_still_rejects_invalid_supplied_values(self):
        base = CleanReview(1, None, None, None, '본문 리뷰', self.now)
        for fields in ({'product_name': ''}, {'review_date': 'bad'}, {'rating': 0}, {'rating': True}):
            with self.subTest(fields=fields), self.assertRaises(ValidationError):
                replace(base, **fields)

    def test_reclean_retains_analysis_for_same_nulls_and_invalidates_changed_metadata(self):
        row = RawReview(source_review_id='1', review_text='사용하기 편리한 리뷰')
        clean = self.clean_rows([row]).reviews[0]
        self.repo.save_analysis(AnalysisResult(clean.id, Sentiment.POSITIVE, .9, self.now, 'fake', 'offline'))
        self.clean.clean_reviews(CleanRequest(CleaningOptions(DuplicatePolicy.UPSERT, 3)))
        self.assertIsNotNone(self.repo.get_review(clean.id).analysis)
        self.repo.save_clean_reviews([replace(clean, rating=5)], DuplicatePolicy.UPSERT)
        self.assertIsNone(self.repo.get_review(clean.id).analysis)


if __name__ == '__main__':
    unittest.main()
